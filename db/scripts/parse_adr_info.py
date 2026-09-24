"""
Clean the ADR Information sheet into one CSV of ADR names.

Usage:
    python3 db/scripts/parse_adr_info.py [--input (location of ADR_Database.xlsx file)] [--output (where you want to put data files)]
    ex: python3 db/scripts/parse_adr_info.py --input db/ADR_Database.xlsx --output db/adr_data
"""

import argparse
import csv
import re
import sys
import warnings
from pathlib import Path

from openpyxl import load_workbook

DB = Path(__file__).resolve().parent.parent
DEFAULT_INPUT = DB / "ADR_Database.xlsx"
DEFAULT_OUTPUT = DB / "adr_data"

SHEET_NAME = "ADR Information"
OUTPUT_NAME = "adr_info.csv"
OUTPUT_COLUMNS = ["ADR Number", "ADR Name"]

ADR_PATTERN = re.compile(r"ADR[-\s]*0*(\d+)", re.IGNORECASE)
EXCEL_ERRORS = {"#N/A", "#REF!", "#VALUE!", "#DIV/0!", "#NAME?", "#NULL!", "#NUM!"}


def clean(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = " ".join(str(value).split())
    return "" if text.upper() in EXCEL_ERRORS else text


def parse_sheet(sheet, warn):
    rows = sheet.iter_rows(values_only=True)
    header = [clean(name) for name in next(rows, None) or []]
    missing = [name for name in OUTPUT_COLUMNS if name not in header]
    if missing:
        raise ValueError(f"{SHEET_NAME}: could not find columns {missing} in header {header}")
    code_index = header.index("ADR Number")
    name_index = header.index("ADR Name")

    cleaned = {}
    for row in rows:
        code = clean(row[code_index]) if code_index < len(row) else ""
        name = clean(row[name_index]) if name_index < len(row) else ""
        if not code:
            continue
        match = ADR_PATTERN.fullmatch(code)
        if not match:
            warn(f"{SHEET_NAME}: unrecognised ADR Number {code!r}")
            continue
        code = f"ADR-{int(match.group(1)):02d}"
        if not name:
            warn(f"{SHEET_NAME}: {code} has no ADR Name")
            continue
        if code in cleaned and cleaned[code] != name:
            warn(f"{SHEET_NAME}: {code} listed twice ({cleaned[code]!r} and {name!r}), keeping the first")
            continue
        cleaned.setdefault(code, name)
    return [{"ADR Number": code, "ADR Name": name} for code, name in cleaned.items()]


def write_csv(path, rows):
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=OUTPUT_COLUMNS)
        writer.writeheader()
        writer.writerows(rows)


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT, help="ADR database workbook")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT, help="directory for the cleaned CSV")
    args = parser.parse_args()

    def warn(message):
        print(f"warning: {message}", file=sys.stderr)

    warnings.filterwarnings(
        "ignore",
        category=UserWarning,
        module="openpyxl",
        message=".*extension is not supported and will be removed",
    )
    workbook = load_workbook(args.input, read_only=True, data_only=True)
    sheet = workbook[SHEET_NAME] if SHEET_NAME in workbook.sheetnames else None
    if sheet is None:
        workbook.close()
        raise SystemExit(f"No sheet named {SHEET_NAME!r} in {args.input}")

    args.output.mkdir(parents=True, exist_ok=True)
    try:
        rows = parse_sheet(sheet, warn)
    finally:
        workbook.close()

    path = args.output / OUTPUT_NAME
    write_csv(path, rows)
    print(f"{SHEET_NAME} -> {path} ({len(rows)} ADRs)")


if __name__ == "__main__":
    main()
