"""
Clean the MASTER Verification sheet into one project-to-ADR CSV.

Usage:
    python3 db/parse_project_to_adr.py [--input (location of ADR_Database.xlsx file)] [--output (where you want to put data files)]
    ex: python3 db/parse_project_to_adr.py --input db/ADR_Database.xlsx --output db/data
"""

import argparse
import csv
import re
import sys
import warnings
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
DEFAULT_INPUT = ROOT / "ADR_Database.xlsx"
DEFAULT_OUTPUT = ROOT / "data"

SHEET_NAME = "MASTER Verification"
OUTPUT_NAME = "project_adr_map.csv"

PROJECT_ID_PATTERN = re.compile(r"SRP8-(\d{1,3})", re.IGNORECASE)
ADR_PATTERN = re.compile(r"ADR-0*(\d+)", re.IGNORECASE)
EXCEL_ERRORS = {"#N/A", "#REF!", "#VALUE!", "#DIV/0!", "#NAME?", "#NULL!", "#NUM!"}
NOT_APPLICABLE = {"N/A", "NA", "-"}


def clean(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    lines = (" ".join(line.split()) for line in str(value).splitlines())
    text = "\n".join(line for line in lines if line)
    return "" if text.upper() in EXCEL_ERRORS else text


def clean_project_id(value, sheet, warn):
    project_id = clean(value)
    if project_id.upper() in NOT_APPLICABLE:
        return ""
    match = PROJECT_ID_PATTERN.fullmatch(project_id)
    if match:
        return f"SRP8-{int(match.group(1)):03d}"
    if project_id:
        warn(f"{sheet}: unrecognised Project ID {project_id!r}")
    return project_id


def clean_adr(value, column, sheet, warn):
    code = clean(value)
    if not code or code.upper() in NOT_APPLICABLE:
        return ""
    match = ADR_PATTERN.fullmatch(code)
    if match:
        return f"ADR-{int(match.group(1)):02d}"
    warn(f"{sheet}: unrecognised {column} {code!r}")
    return code


def header_names(header):
    names = []
    for index, name in enumerate(header):
        names.append(clean(name) or f"Column {index + 1}")
    while names and re.fullmatch(r"Column \d+", names[-1]):
        names.pop()
    if "Project ID" not in names:
        raise ValueError(f"{SHEET_NAME}: could not find a Project ID column in header {names}")
    return names


def parse_sheet(sheet, warn):
    rows = sheet.iter_rows(values_only=True)
    header = next(rows, None)
    if header is None:
        raise ValueError(f"{SHEET_NAME}: sheet is empty")
    columns = header_names(header)
    adr_columns = [name for name in columns if name.endswith("ADR")]

    cleaned = []
    for row in rows:
        record = {
            name: clean(row[index]) if index < len(row) else ""
            for index, name in enumerate(columns)
        }
        if not any(record.values()):
            continue
        record["Project ID"] = clean_project_id(record["Project ID"], sheet.title, warn)
        for name in adr_columns:
            record[name] = clean_adr(record[name], name, sheet.title, warn)
        cleaned.append(record)
    return columns, cleaned


def write_csv(path, columns, rows):
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns)
        writer.writeheader()
        writer.writerows(rows)


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT, help="ADR database workbook")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT, help="directory for the cleaned CSV")
    args = parser.parse_args()

    def warn(message):
        print(f"warning: {message}", file=sys.stderr)

    warnings.filterwarnings(
        "ignore",
        category=UserWarning,
        module="openpyxl",
        message=".*extension is not supported and will be removed",
    )
    workbook = load_workbook(args.input, read_only=True, data_only=True)
    sheet = workbook[SHEET_NAME] if SHEET_NAME in workbook.sheetnames else None
    if sheet is None:
        workbook.close()
        raise SystemExit(f"No sheet named {SHEET_NAME!r} in {args.input}")

    args.output.mkdir(parents=True, exist_ok=True)
    try:
        columns, rows = parse_sheet(sheet, warn)
    finally:
        workbook.close()

    path = args.output / OUTPUT_NAME
    write_csv(path, columns, rows)
    linked = sum(1 for row in rows if any(row[name] for name in columns if name.endswith("ADR")))
    print(f"{SHEET_NAME} -> {path} ({len(rows)} projects, {linked} with an ADR)")


if __name__ == "__main__":
    main()
